# -*- coding: utf-8 -*-
"""Dung Danh muc Yeu cau Nghiep vu (BRL) theo dung cau truc file BR mau JVC Eco.

File nay la TAI LIEU NOI BO: dung de boc gio tung cap Junior/Senior/Expert va biet gia von.
KHONG gui cho khach. Ban gui khach la phuong an HTML (buoc 4), chi ghi tong tien phan phat trien.

Dung:
    import vbr
    vbr.dung(rows, "…/Danh mục Yêu cầu Nghiệp vụ Phần mềm - <Khách>.xlsx",
             khach="Công ty TNHH ABC", to_chuc=["Công ty TNHH ABC"], diem_xac_nhan=XN)

Moi phan tu `rows` la mot dict:
    id   ma yeu cau, vi du "BR-MH-01" (hoac "BR-001" neu danh so phang)
    ten  mo ta yeu cau - mot cau ngan
    ct   mo ta chi tiet - nguyen van y khach, khong dien rong
    dep  phu thuoc (chuoi, co the de rong)
    nv   phan loai theo nghiep vu - lay tu sheet "Phân loại Nghiệp vụ"
    nhom phan loai theo Nhom - lay tu cung sheet (cot Phan loai)
    pbd  phong ban de xuat   | pbs  phong ban su dung - lay tu sheet "Phòng ban"
    pl   "Tiêu chuẩn" | "Cấu hình" | "Tích hợp" | "Tùy chỉnh" | "Phát triển" | "Cần làm rõ" | "Loại bỏ"
    gp   giai phap: pham vi + cach dap ung + co so ma nguon (model/field/module that)
    note ghi chu (tuy chon, noi vao cuoi mo ta chi tiet)
    gio  dict gio cong theo giai doan, chi dat cho hang muc phai lam them:
         {"ks": (gio, cap), "tk": …, "lt": …, "qa": …, "dt": …, "ch": …, "tl": …, "gs": …}
         cap la "Junior" | "Senior" | "Expert". Bao hanh (bh) tu tinh = 15% gio lap trinh.
"""
from __future__ import annotations

import copy
import os
import re
import sys
import warnings

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.formula.translate import Translator
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings("ignore")

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MAU = os.path.join(KIT, "mau", "danh-muc-br.xlsx")

RBN = "R - Business Requirements"
NC = 40                                   # cot A..AN
# cot gio cua tung giai doan; cot +1 la cap bac, cot +2 la chi phi (cong thuc co san trong mau)
COT_GIO = {"ks": 10, "tk": 13, "lt": 16, "qa": 19, "dt": 22, "ch": 25, "bh": 28, "tl": 31, "gs": 34}
TEN_GIAI_DOAN = {
    "ks": "Khảo sát", "tk": "Phân tích & Thiết kế", "lt": "Lập trình & Phát triển",
    "qa": "QA & Kiểm thử thủ công", "dt": "Đào tạo và Kiểm thử chấp thuận",
    "ch": "Cấu hình & Triển khai", "bh": "Bảo hành (Bug fix)",
    "tl": "Tài liệu / video hướng dẫn sử dụng", "gs": "Giám sát hậu triển khai",
}
COT_CHI_PHI = {"ks": "L", "tk": "O", "lt": "R", "qa": "U", "dt": "X", "ch": "AA", "bh": "AD", "tl": "AG", "gs": "AJ"}
PHAN_LOAI = ["Tiêu chuẩn", "Cấu hình", "Tích hợp", "Tùy chỉnh", "Phát triển", "Cần làm rõ", "Loại bỏ"]

# He so "tan dung AI toi da" cho hang muc phat trien them: buoc AI lam duoc chi tinh phan
# chuyen vien duyet, sua, tich hop. Phan tich thiet ke va dao tao/nghiem thu voi khach giu nguyen.
HE_SO_AI = {"lt": 0.4, "qa": 0.5, "tl": 0.3}


def ap_he_so_ai(gio: dict) -> dict:
    """Nhan he so AI vao gio goc. Goi DUNG MOT LAN tren gio chua nhan."""
    return {k: ((round(v[0] * HE_SO_AI.get(k, 1), 2)), v[1]) for k, v in gio.items()}


def _sach(duong_dan: str) -> str:
    """Lam sach TEN TEP (giu nguyen thu muc) - bo ky tu khong hop le tren he tep."""
    thu_muc, ten = os.path.split(duong_dan)
    return os.path.join(thu_muc, re.sub(r'[\\/:*?"<>|]', "-", ten))


def dung(rows, out: str, khach: str, to_chuc=None, diem_xac_nhan=None,
         ghi_chu_chi_phi: str = "", mau: str = MAU, sap_xep=None) -> str:
    """Dung workbook BR va ghi ra `out`. Tra ve duong dan da ghi."""
    wb = openpyxl.load_workbook(mau)
    ws = wb[RBN]
    RB = f"'{RBN}'!"

    st = {c: copy.copy(ws.cell(2, c)._style) for c in range(1, NC + 1)}
    st_tot = {c: copy.copy(ws.cell(3, c)._style) for c in range(1, NC + 1)}
    tpl_f = {c: ws.cell(2, c).value for c in range(1, NC + 1)}
    tot_f = {c: ws.cell(3, c).value for c in range(1, NC + 1)}
    ws.delete_rows(2, 2)                                   # bo dong mau + dong Tong, dung lai tu dau

    if sap_xep:
        rows = sorted(rows, key=sap_xep)

    for i, x in enumerate(rows):
        r = i + 2
        for c in range(1, NC + 1):
            ws.cell(r, c)._style = copy.copy(st[c])
            f = tpl_f[c]
            if isinstance(f, str) and f.startswith("="):
                ws.cell(r, c).value = Translator(f, origin=f"{ws.cell(1, c).column_letter}2").translate_formula(
                    ws.cell(r, c).coordinate)
        ws.cell(r, 1).value = "=ROW()-1"
        ct = x.get("ct", "")
        if x.get("note"):
            ct = (ct + "\n" if ct else "") + "Ghi chú: " + x["note"]
        for c, v in {2: x["id"], 3: x.get("ten"), 4: ct or None, 5: x.get("dep") or None,
                     6: x.get("nv"), 7: x.get("nhom"), 8: x.get("pbd"), 9: x.get("pbs"),
                     39: x.get("pl"), 40: x.get("gp")}.items():
            ws.cell(r, c).value = v
        for k, (h, cap) in (x.get("gio") or {}).items():
            if k == "bh":
                continue
            ws.cell(r, COT_GIO[k]).value = h
            ws.cell(r, COT_GIO[k] + 1).value = cap
        if (x.get("gio") or {}).get("lt"):                 # bao hanh = 15% gio lap trinh
            ws.cell(r, COT_GIO["bh"]).value = f"=P{r}*15%"
            ws.cell(r, COT_GIO["bh"] + 1).value = x["gio"].get("bh", (0, x["gio"]["lt"][1]))[1]

    cuoi = len(rows) + 1
    tong = cuoi + 1
    for c in range(1, NC + 1):
        ws.cell(tong, c)._style = copy.copy(st_tot[c])
        v = tot_f[c]
        if isinstance(v, str) and v.startswith("="):
            ws.cell(tong, c).value = v
    ws.cell(tong, 2).value = "Tổng"

    t = ws.tables["Danh_mục_Yêu_cầu_nghiệp_vụ"]
    t.ref = f"A1:AN{tong}"
    t.autoFilter.ref = f"A1:AN{cuoi}"
    t.autoFilter.filterColumn = []
    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"' + ",".join(PHAN_LOAI) + '"', allow_blank=True)
    dv.add(f"AM2:AM{cuoi}")
    ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1="'Bảng giá theo manday'!$A$2:$A$4", allow_blank=True)
    for col in ["K", "N", "Q", "T", "W", "Z", "AC", "AF", "AI"]:
        dv2.add(f"{col}2:{col}{cuoi}")
    ws.add_data_validation(dv2)

    _dashboard(wb, rows, cuoi, RB, ghi_chu_chi_phi)
    _to_chuc(wb, to_chuc or [khach])
    if diem_xac_nhan:
        _xac_nhan(wb, diem_xac_nhan, khach)

    wb.properties.creator = "Viindoo"
    wb.properties.title = f"Danh mục Yêu cầu Nghiệp vụ Phần mềm - {khach}"
    wb.active = wb.sheetnames.index("Dashboard (Chart)")
    out = _sach(out)
    wb.save(out)
    return out


# ----------------------------------------------------------------- dashboard
def _dashboard(wb, rows, cuoi, RB, ghi_chu):
    dash = wb["Dashboard"]
    dash._pivots = []
    for row in dash.iter_rows():
        for c in row:
            c.value = None

    def rng(L):
        return f"{RB}${L}$2:${L}${cuoi}"

    def uniq(key):
        out = []
        for x in rows:
            v = x.get(key)
            if v and v not in out:
                out.append(v)
        return out

    def khoi(col, tieu_de, key, L, cong_thuc):
        dash.cell(2, col).value = tieu_de
        dash.cell(3, col).value = tieu_de
        dash.cell(3, col + 1).value = "Chi phí" if cong_thuc == "sum" else "Số lượng yêu cầu"
        vals = uniq(key)
        for i, v in enumerate(vals):
            o = dash.cell(4 + i, col)
            o.value = v
            if cong_thuc == "sum":
                dash.cell(4 + i, col + 1).value = f'=SUMIFS({rng("AL")},{rng(L)},{o.coordinate})'
            else:
                dash.cell(4 + i, col + 1).value = f"=COUNTIFS({rng(L)},{o.coordinate})"
        return 4, 3 + max(len(vals), 1)

    blocks = {
        "nhom": (1, khoi(1, "Phân loại theo Nhóm", "nhom", "G", "count")),
        "nv": (4, khoi(4, "Phân loại theo nghiệp vụ", "nv", "F", "count")),
        "pbd": (7, khoi(7, "Phòng ban đề xuất", "pbd", "H", "count")),
        "pbs": (10, khoi(10, "Phòng ban sử dụng", "pbs", "I", "count")),
    }
    dash.cell(2, 13).value = "Số ngày công theo từng cấp bậc"
    dash.cell(3, 13).value = "Cấp bậc"
    dash.cell(3, 14).value = "Số lượng ngày công"
    for i, cap in enumerate(["Expert", "Senior", "Junior"]):
        r = 4 + i
        dash.cell(r, 13).value = cap
        phan = [f"SUMIFS({rng(openpyxl.utils.get_column_letter(c))},"
                f"{rng(openpyxl.utils.get_column_letter(c + 1))},M{r})" for c in COT_GIO.values()]
        dash.cell(r, 14).value = "=(" + "+".join(phan) + ")/8"
    blocks["cap"] = (13, (4, 6))
    blocks["c_nhom"] = (16, khoi(16, "Chi phí theo Nhóm", "nhom", "G", "sum"))
    blocks["c_nv"] = (19, khoi(19, "Chi phí theo nghiệp vụ", "nv", "F", "sum"))
    blocks["c_pbs"] = (22, khoi(22, "Chi phí theo Phòng ban sử dụng", "pbs", "I", "sum"))
    dash.cell(2, 25).value = "Chi phí theo cấp bậc"
    dash.cell(3, 25).value = "Cấp bậc"
    dash.cell(3, 26).value = "Chi phí"
    for i, (cap, o) in enumerate([("Expert", "B4"), ("Senior", "B3"), ("Junior", "B2")]):
        dash.cell(4 + i, 25).value = cap
        dash.cell(4 + i, 26).value = f"=N{4 + i}*'Bảng giá theo manday'!{o}"
    blocks["c_cap"] = (25, (4, 6))
    dash.cell(2, 28).value = "Chi phí theo cơ cấu"
    dash.cell(3, 28).value = "Cơ cấu chi phí"
    dash.cell(3, 29).value = "Chi phí"
    for i, k in enumerate(COT_GIO):
        dash.cell(4 + i, 28).value = TEN_GIAI_DOAN[k]
        dash.cell(4 + i, 29).value = f"=SUM({rng(COT_CHI_PHI[k])})"
    blocks["c_cc"] = (28, (4, 3 + len(COT_GIO)))

    ch = wb["Dashboard (Chart)"]
    ch._charts = []
    ch["B4"] = "DASHBOARD SỐ LƯỢNG"
    ch["B38"] = "DASHBOARD CHI PHÍ"
    if ghi_chu:
        ch["B39"] = ghi_chu
    bieu_do = [
        ("nhom", "Phân bố yêu cầu theo Nhóm chức năng", "B6"), ("nv", "Phân bố yêu cầu theo Nghiệp vụ", "J6"),
        ("pbd", "Số yêu cầu do các Phòng ban đề xuất", "B21"), ("pbs", "Số yêu cầu gắn với Phòng ban sử dụng", "J21"),
        ("cap", "Tổng số ngày công theo Cấp bậc nhân sự", "R6"),
        ("c_nhom", "Tổng chi phí theo Nhóm chức năng", "B41"), ("c_nv", "Tổng chi phí theo Nghiệp vụ", "J41"),
        ("c_pbs", "Chi phí phân bổ theo Phòng ban sử dụng", "B56"), ("c_cap", "Tổng chi phí theo Cấp bậc nhân sự", "J56"),
        ("c_cc", "Cơ cấu chi phí theo Hạng mục công việc", "R41"),
    ]
    for key, tieu_de, neo in bieu_do:
        col, (r0, r1) = blocks[key]
        bc = BarChart()
        bc.type = "bar"
        bc.title = tieu_de
        bc.legend = None
        bc.add_data(Reference(dash, min_col=col + 1, min_row=3, max_row=r1), titles_from_data=True)
        bc.set_categories(Reference(dash, min_col=col, min_row=r0, max_row=r1))
        bc.height, bc.width = 7.2, 14
        ch.add_chart(bc, neo)


def _to_chuc(wb, ten_list):
    oc = wb["Tổ chức  Công ty"]
    st = [copy.copy(oc.cell(2, c)._style) for c in (1, 2)]
    for i, n in enumerate(ten_list):
        oc.cell(2 + i, 1).value = "=ROW()-1"
        oc.cell(2 + i, 2).value = n
        oc.cell(2 + i, 1)._style = copy.copy(st[0])
        oc.cell(2 + i, 2)._style = copy.copy(st[1])


def _xac_nhan(wb, diem, khach):
    """Sheet 'Điểm cần xác nhận' - cung noi dung voi muc xac nhan trong tai lieu HTML."""
    import vykien
    ws = wb.create_sheet("Điểm cần xác nhận", 0)
    ws["A1"] = ("Mỗi điểm dưới đây Viindoo đã nêu cách hiểu và đề xuất của mình kèm căn cứ. "
                f"{khach} tích (x) vào dòng phù hợp hoặc ghi ý kiến ở cột cuối.")
    ws["A1"].alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A1:F1")
    cot = ["Mã", "Điểm cần xác nhận", "Viindoo hiểu và đề xuất", "Nội dung cần xác nhận", "Tích (x)",
           f"Ý kiến của {khach}"]
    for c, v in enumerate(cot, start=1):
        o = ws.cell(2, c)
        o.value = v
        o.font = openpyxl.styles.Font(bold=True, color="FFFFFF")
        o.fill = openpyxl.styles.PatternFill("solid", fgColor="00BBCE")
        o.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="center")
    for i, (ma, ten, hieu, tich) in enumerate(vykien.dong_xlsx(diem), start=3):
        for c, v in enumerate([ma, ten, hieu, tich, "", ""], start=1):
            o = ws.cell(i, c)
            o.value = v
            o.alignment = openpyxl.styles.Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", [10, 30, 70, 46, 9, 40]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"
